"""
AD 使用者 / 群組查詢（LDAP 直查）
====================================
直接連 Domain Controller 查詢 AD 屬性。

用法：
    # 使用者基本資訊
    python fetch_ad_info.py --account JIAHUA.WU

    # 使用者所屬群組
    python fetch_ad_info.py --account JIAHUA.WU --groups

    # 查詢單一群組成員
    python fetch_ad_info.py --group G-Delta-rollout_admin

    # 列出所有群組（可加 pattern 過濾）
    python fetch_ad_info.py --all-groups
    python fetch_ad_info.py --all-groups --filter "G-Delta*"

    # 批次查詢工單帳號
    python fetch_ad_info.py --input delta_tickets_clean.json --output ad_info.json

    # 輸出成 Excel（群組清單 + 群組成員 1對多）
    python fetch_ad_info.py --all-groups --filter "G-Delta*" --output groups.xlsx

環境變數（選填，不設則用目前 Kerberos ticket）：
    AD_DC, AD_BASE_DN, AD_USER, AD_PASSWORD
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

# Windows cp950 terminal 遇到外籍人名會炸，強制 utf-8 輸出
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    from ldap3 import Server, Connection, SASL, GSSAPI, NTLM, SUBTREE, ALL_ATTRIBUTES
except ImportError:
    raise RuntimeError("請先安裝 ldap3：pip install ldap3")

DC_HOST = os.getenv("AD_DC",      "TWTPEDCS02")
BASE_DN = os.getenv("AD_BASE_DN", "DC=delta,DC=corp")
AD_USER = os.getenv("AD_USER",    "")
AD_PASS = os.getenv("AD_PASSWORD","")

ATTR_MAP = {
    "sAMAccountName":       "AD Account",
    "cn":                   "CN",
    "mail":                 "Mail",
    "department":           "部門",
    "manager":              "主管",
    "extensionAttribute15": "工號",
    "extensionAttribute1":  "BG",
    "extensionAttribute2":  "BU",
}
USER_ATTRS  = list(ATTR_MAP.keys())
GROUP_ATTRS = ["cn", "description", "managedBy", "member"]


def get_connection():
    server = Server(DC_HOST, get_info=None)
    if AD_USER and AD_PASS:
        return Connection(server, user=AD_USER, password=AD_PASS, authentication=NTLM, auto_bind=True)
    return Connection(server, authentication=SASL, sasl_mechanism=GSSAPI, auto_bind=True)


def _cn_from_dn(dn: str) -> str:
    m = re.search(r"CN=([^,]+)", dn or "")
    return m.group(1) if m else dn


def fetch_user(conn, ad_account: str, include_groups=False) -> dict:
    attrs = USER_ATTRS + (["memberOf"] if include_groups else [])
    conn.search(BASE_DN, f"(sAMAccountName={ad_account})", search_scope=SUBTREE, attributes=attrs)
    if not conn.entries:
        return {"AD Account": ad_account}

    entry = conn.entries[0]
    result = {}
    for attr, label in ATTR_MAP.items():
        raw = getattr(entry, attr, None)
        val = str(raw.value) if (raw and raw.value) else ""
        if label == "主管" and "CN=" in val:
            val = _cn_from_dn(val)
        if label in ("BU", "BG") and "/" in val:
            val = val.split("/")[0]
        result[label] = val

    if include_groups:
        raw = getattr(entry, "memberOf", None)
        values = raw.value if (raw and raw.value) else []
        if isinstance(values, str):
            values = [values]
        result["Groups"] = [_cn_from_dn(dn) for dn in values]

    return result


def _resolve_by_memberof(conn, group_dn: str) -> list:
    """用 memberOf=<group_dn> 反查所有成員（含 BU/BG），避免大型 OR filter。"""
    conn.search(
        BASE_DN,
        f"(memberOf={group_dn})",
        search_scope=SUBTREE,
        attributes=["sAMAccountName", "cn", "mail", "extensionAttribute1", "extensionAttribute2"],
    )
    members = []
    for e in conn.entries:
        def _ext(attr, _e=e):
            raw = getattr(_e, attr, None)
            val = str(raw.value) if (raw and raw.value) else ""
            return val.split("/")[0] if "/" in val else val
        members.append({
            "AD Account": str(e.sAMAccountName.value) if e.sAMAccountName.value else "",
            "CN":         str(e.cn.value)             if e.cn.value             else "",
            "Mail":       str(e.mail.value)           if e.mail.value           else "",
            "BG":         _ext("extensionAttribute1"),
            "BU":         _ext("extensionAttribute2"),
        })
    members.sort(key=lambda x: x["AD Account"])
    return members


def fetch_group(conn, group_name: str) -> dict:
    conn.search(BASE_DN, f"(cn={group_name})", search_scope=SUBTREE, attributes=GROUP_ATTRS)
    if not conn.entries:
        return {"Group": group_name, "Members": [], "error": "查無此群組"}

    entry = conn.entries[0]
    group_dn = entry.entry_dn

    managed_by = ""
    raw_mgr = getattr(entry, "managedBy", None)
    if raw_mgr and raw_mgr.value:
        managed_by = _cn_from_dn(str(raw_mgr.value))

    members = _resolve_by_memberof(conn, group_dn)
    return {
        "Group":       str(entry.cn.value) if entry.cn.value else group_name,
        "Description": str(entry.description.value) if getattr(entry, "description", None) and entry.description.value else "",
        "ManagedBy":   managed_by,
        "MemberCount": len(members),
        "Members":     members,
    }


def fetch_all_groups(conn, pattern: str = "*") -> list:
    """列出所有符合 pattern 的群組（分頁查詢，突破 AD 1000 筆上限）。"""
    entries = conn.extend.standard.paged_search(
        BASE_DN,
        f"(&(objectClass=group)(cn={pattern}))",
        search_scope=SUBTREE,
        attributes=["cn", "description", "managedBy", "member"],
        paged_size=500,
        generator=False,
    )
    groups = []
    for entry in entries:
        if entry.get("type") != "searchResEntry":
            continue
        attrs = entry.get("attributes", {})
        members = attrs.get("member", [])
        cnt = len(members) if isinstance(members, list) else (1 if members else 0)
        mgr_raw = attrs.get("managedBy", "")
        if isinstance(mgr_raw, list):
            mgr_raw = mgr_raw[0] if mgr_raw else ""
        managed_by = _cn_from_dn(str(mgr_raw)) if mgr_raw else ""
        cn_val = attrs.get("cn", "")
        if isinstance(cn_val, list):
            cn_val = cn_val[0] if cn_val else ""
        desc = attrs.get("description", "")
        if isinstance(desc, list):
            desc = desc[0] if desc else ""
        groups.append({
            "Group":       str(cn_val),
            "Description": str(desc),
            "ManagedBy":   managed_by,
            "MemberCount": cnt,
            "_dn":         entry.get("dn", ""),
        })
    groups.sort(key=lambda x: x["Group"])
    return groups


def save_xlsx(groups_with_members: list, output_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    header_fill = PatternFill(start_color="FFEBF4F4", end_color="FFEBF4F4", fill_type="solid")

    def _apply_header(ws, cols):
        ws.append(cols)
        for col in range(1, len(cols) + 1):
            ws.cell(1, col).fill = header_fill

    def _auto_width(ws):
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(w + 4, 60)

    # ── Sheet 1：群組清單 ──────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "群組清單"
    _apply_header(ws1, ["群組名稱", "說明", "管理者", "成員數", "BG"])
    for g in groups_with_members:
        bgs = sorted({m.get("BG", "") for m in g.get("Members", []) if m.get("BG")})
        ws1.append([g["Group"], g.get("Description", ""), g.get("ManagedBy", ""), g["MemberCount"], ", ".join(bgs)])
    _auto_width(ws1)

    # ── 依 BG 分 tab ──────────────────────────────────────────────────────
    # 先收集每個 BG 的所有 (群組, 成員) row
    bg_rows: dict[str, list] = {}
    for g in groups_with_members:
        for m in g.get("Members", []):
            bg = m.get("BG", "") or "其他"
            bg_rows.setdefault(bg, []).append(
                [g["Group"], m["AD Account"], m["CN"], m["Mail"], m.get("BU", "")]
            )

    for bg in sorted(bg_rows.keys()):
        # Excel sheet 名稱最長 31 字元，去掉非法字元
        sheet_name = re.sub(r'[\\/*?:\[\]]', "_", bg)[:31]
        ws = wb.create_sheet(sheet_name)
        _apply_header(ws, ["群組", "AD Account", "CN", "Mail", "BU"])
        for row in bg_rows[bg]:
            ws.append(row)
        _auto_width(ws)

    wb.save(output_path)
    total = sum(len(v) for v in bg_rows.values())
    print(f"[OK] 已儲存 {output_path}（群組清單 + {len(bg_rows)} 個 BG tab，共 {total} 筆成員）")


def fetch_batch(conn, accounts: list) -> list:
    results = []
    seen = set()
    for acc in accounts:
        if not acc or acc in seen:
            continue
        seen.add(acc)
        print(f"  {acc:25s}", end=" ", flush=True)
        info = fetch_user(conn, acc)
        print(f"BU={info.get('BU','?'):10s}  BG={info.get('BG','?'):6s}  {info.get('部門','')}")
        results.append(info)
    return results


def main():
    parser = argparse.ArgumentParser(description="查詢 AD 使用者 / 群組（LDAP）")
    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--account",    help="單一使用者帳號")
    src.add_argument("--group",      help="查詢單一群組成員")
    src.add_argument("--all-groups", action="store_true", help="列出所有群組")
    src.add_argument("--input",      help="工單 JSON，批次查詢帳號")
    parser.add_argument("--filter",    default="*", help="--all-groups 的 CN pattern，例如 G-Delta*")
    parser.add_argument("--groups",    action="store_true", help="（配合 --account）列出所屬群組")
    parser.add_argument("--members",   action="store_true", help="（配合 --all-groups）同時展開成員（輸出 xlsx 自動開啟）")
    parser.add_argument("--all-attrs", action="store_true", help="（配合 --account）印出所有屬性")
    parser.add_argument("--output",    help="輸出 JSON 或 xlsx 檔")
    args = parser.parse_args()

    print(f"[*] 連線至 {DC_HOST}...")
    conn = get_connection()
    print(f"[*] 連線成功\n")

    results = None

    # ── 使用者 ────────────────────────────────────────────────────────────────
    if args.account:
        if args.all_attrs:
            conn.search(BASE_DN, f"(sAMAccountName={args.account})", search_scope=SUBTREE, attributes=ALL_ATTRIBUTES)
            if conn.entries:
                for attr in sorted(conn.entries[0].entry_attributes):
                    val = conn.entries[0][attr].value
                    if val:
                        print(f"  {attr:40s}: {val}")
            else:
                print("  (查無此帳號)")
            conn.unbind()
            return

        info = fetch_user(conn, args.account, include_groups=args.groups)
        for k, v in info.items():
            if k == "Groups":
                print(f"\n  所屬群組（{len(v)} 個）：")
                for g in v:
                    print(f"    - {g}")
            else:
                print(f"  {k:10s}: {v}")
        results = [info]

    # ── 單一群組 ──────────────────────────────────────────────────────────────
    elif args.group:
        info = fetch_group(conn, args.group)
        print(f"  群組      : {info['Group']}")
        print(f"  說明      : {info.get('Description','')}")
        print(f"  管理者    : {info.get('ManagedBy','')}")
        print(f"  成員數    : {info['MemberCount']}\n")
        for m in info["Members"]:
            print(f"  {m['AD Account']:25s}  {m['CN']:30s}  {m['Mail']}")
        results = [info]

        # 單一群組也支援 --output（不需等到最下方共用邏輯）
        if args.output:
            if args.output.endswith(".xlsx"):
                save_xlsx(results, args.output)
            else:
                Path(args.output).write_text(
                    json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                print(f"[OK] 已輸出至 {args.output}")
            conn.unbind()
            return

    # ── 所有群組 ──────────────────────────────────────────────────────────────
    elif args.all_groups:
        groups = fetch_all_groups(conn, args.filter)
        print(f"  {'群組名稱':40s}  {'成員數':>5}  管理者")
        print("  " + "-"*70)
        for g in groups:
            print(f"  {g['Group']:40s}  {g['MemberCount']:>5}  {g['ManagedBy']}")
        print(f"\n  共 {len(groups)} 個群組")

        if args.members or (args.output and args.output.endswith(".xlsx")):
            print(f"\n[*] 展開成員中（共 {len(groups)} 個群組）...")
            for g in groups:
                print(f"  {g['Group']}...", flush=True)
                dn = g.pop("_dn", "")
                if dn:
                    g["Members"] = _resolve_by_memberof(conn, dn)
                    g["MemberCount"] = len(g["Members"])
                else:
                    g["Members"] = []
        results = groups

    # ── 批次 ──────────────────────────────────────────────────────────────────
    else:
        path = Path(args.input)
        if not path.exists():
            print(f"[!] 找不到 {args.input}")
            conn.unbind()
            return
        tickets = json.loads(path.read_text(encoding="utf-8"))
        accounts = [t.get("requester", "").split()[0] for t in tickets if t.get("requester")]
        print(f"[*] 共 {len(set(accounts))} 個不重複帳號\n")
        results = fetch_batch(conn, accounts)

    # ── 輸出 ──────────────────────────────────────────────────────────────────
    if args.output and results:
        if args.output.endswith(".xlsx"):
            save_xlsx(results, args.output)
        else:
            Path(args.output).write_text(
                json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            print(f"[OK] 已輸出至 {args.output}")

    conn.unbind()


if __name__ == "__main__":
    main()
